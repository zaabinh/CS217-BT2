"""Create the printable assignment workbook using openpyxl 3.1+."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Tai_lieu_nop_bai" / "Bang_phan_cong_cong_viec.xlsx"
MEMBERS = [
    ("Đoàn Hữu Gia Bình", "24520192", "Lập trình, kiểm thử và hoàn thiện Bài 1"),
    ("Nguyễn Thái Bảo", "24520173", "Xây dựng tài liệu và báo cáo Bài 1"),
    ("Bùi Huy Chương", "24520230", "Lập trình, kiểm thử và hoàn thiện Bài 2"),
    ("Võ Hoài Chiều", "24520220", "Xây dựng tài liệu và báo cáo Bài 2"),
]


def title(label: str, text: str) -> CellRichText:
    return CellRichText([
        TextBlock(InlineFont(rFont="Times New Roman", sz=13, b=True, color="000000"), label),
        TextBlock(InlineFont(rFont="Times New Roman", sz=13, color="000000"), text),
    ])


def main() -> None:
    """Write and reopen the workbook to verify content, styles and print settings."""
    workbook = Workbook()
    workbook.properties.title = "Bảng phân công công việc - Bài tập 2"
    workbook.properties.subject = "Các hệ cơ sở tri thức"
    sheet = workbook.active
    sheet.title = "Phân công công việc"
    sheet.sheet_view.showGridLines = False
    edge = Side(style="thin", color="000000")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for row in sheet.iter_rows(min_row=1, max_row=9, min_col=1, max_col=3):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=13, color="000000")
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in range(1, 5):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(row, 1).font = Font(name="Times New Roman", size=13, color="000000", bold=True)
    sheet["A1"] = title("Môn:", " Các hệ cơ sở tri thức")
    sheet["A2"] = title("Giảng viên:", " PGS.TS Nguyễn Đình Hiển")
    sheet["A3"] = "BÀI TẬP 2 - NHÓM"
    sheet["A4"] = "Bảng phân công công việc"
    for column, label in enumerate(("Tên", "MSSV", "Nội dung công việc"), 1):
        cell = sheet.cell(5, column, label)
        cell.font = Font(name="Times New Roman", size=13, color="000000", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, member in enumerate(MEMBERS, 6):
        for col, value in enumerate(member, 1):
            sheet.cell(row, col, value)
        sheet.cell(row, 2).number_format = "@"
        sheet.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, width in (("A", 27), ("B", 15), ("C", 59)):
        sheet.column_dimensions[column].width = width
    for row, height in ((1, 26), (2, 29), (3, 30), (4, 30), (5, 28)):
        sheet.row_dimensions[row].height = height
    for row in range(6, 10):
        sheet.row_dimensions[row].height = 48
    sheet.freeze_panes = "A6"
    sheet.print_area = "A1:C9"
    sheet.print_title_rows = "1:5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_options.horizontalCentered = True
    sheet.page_margins = PageMargins(left=0.35, right=0.35, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    verify(OUTPUT)
    print(OUTPUT)


def verify(path: Path) -> None:
    """Assert the persisted workbook satisfies the requested submission format."""
    sheet = load_workbook(path, rich_text=True).active
    assert {str(r) for r in sheet.merged_cells.ranges} == {f"A{i}:C{i}" for i in range(1, 5)}
    assert [str(sheet.cell(i, 1).value) for i in range(1, 5)] == [
        "Môn: Các hệ cơ sở tri thức", "Giảng viên: PGS.TS Nguyễn Đình Hiển",
        "BÀI TẬP 2 - NHÓM", "Bảng phân công công việc",
    ]
    assert list(sheet.iter_rows(min_row=6, max_row=9, values_only=True)) == MEMBERS
    for row in sheet:
        for cell in row:
            if cell.value is None:
                continue  # Merged-cell placeholders have no independent content.
            assert cell.font.name == "Times New Roman" and cell.font.sz == 13
            assert cell.font.color.rgb[-6:] == "000000"
            assert all(getattr(cell.border, side).style == "thin" for side in ("left", "right", "top", "bottom"))
            assert cell.alignment.wrap_text and cell.alignment.vertical == "center"
    for row in (1, 2):
        value = sheet.cell(row, 1).value
        assert value[0].font.b and not value[1].font.b
        assert all(block.font.rFont == "Times New Roman" and block.font.sz == 13
                   and block.font.color.rgb[-6:] == "000000" for block in value)
    for cell in sheet[5]:
        assert cell.font.b and cell.fill.fgColor.rgb[-6:] == "D9D9D9"
    assert all(sheet.cell(row, 2).alignment.horizontal == "center" for row in range(6, 10))
    assert str(sheet.page_setup.paperSize) == str(sheet.PAPERSIZE_A4)
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 1


if __name__ == "__main__":
    main()
